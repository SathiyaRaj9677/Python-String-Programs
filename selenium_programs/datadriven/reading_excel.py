
import os
import openpyxl

# File --> Workbook --> Sheets --> Rows --> Cells

filename = "data.xlsx"
search_path = os.path.expanduser("~")

# Find the Excel file
file_path = None

for root, dirs, files in os.walk(search_path):
    if filename in files:
        file_path = os.path.join(root, filename)
        break

if file_path is None:
    print("Excel file not found!")
else:
    print("File found:", file_path)

    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook["Sheet1"]

    rows = sheet.max_row
    cols = sheet.max_column

    # Read all rows and columns from the Excel sheet
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            print(sheet.cell(r, c).value, end="       ")
        print()

    workbook.close()
